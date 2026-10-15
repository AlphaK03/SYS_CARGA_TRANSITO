"""Revision acotada: datos ficticios, red bloqueada y Excel simulado.

Los expectedFailure documentan bugs reproducidos, pendientes de correccion.
No son comprobaciones satisfactorias del comportamiento de la aplicacion.
"""
import base64
from datetime import datetime
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from database.client_database import ClientDatabase
from email_service.email_sender import EmailSender
from gui.app import Application
from processing.file_processor import FileProcessor, _save_excel_with_format
from utils.config_manager import ConfigManager


class RevisionDevTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="astra_revision_")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        for target in ("socket.socket.connect", "socket.socket.connect_ex",
                       "requests.sessions.Session.request",
                       "keyring.get_password", "keyring.set_password",
                       "keyring.delete_password", "win32com.client.DispatchEx"):
            blocker = patch(target, side_effect=AssertionError("Acceso real bloqueado"))
            blocker.start()
            self.addCleanup(blocker.stop)

    def export_report(self, frame, with_table=False):
        template = self.folder / "template.xlsx"
        output = self.folder / "report.xlsx"
        workbook = Workbook()
        if with_table:
            worksheet = workbook.active
            worksheet["A4"] = "Consignee"
            worksheet["A5"] = "Cliente ficticio"
            worksheet.add_table(Table(displayName="OldTable", ref="A4:A5"))
        workbook.save(template)
        workbook.close()
        _save_excel_with_format(frame, str(template), str(output))
        result = load_workbook(output)
        self.addCleanup(result.close)
        return result.active

    def make_sender(self, token="fake-token"):
        sender = EmailSender.__new__(EmailSender)
        sender.msal_app = Mock()
        sender.msal_app.get_accounts.return_value = [{"username": "test@example.invalid"}]
        sender.msal_app.acquire_token_silent.return_value = (
            {"access_token": token} if token else None
        )
        attachment = self.folder / "attachment.txt"
        attachment.write_bytes(b"reporte ficticio")
        return sender, str(attachment)

    def test_report_dates_and_internal_columns(self):
        sheet = self.export_report(pd.DataFrame({
            "Consignee": ["Cliente ficticio"],
            "ETA": ["9/8/25"],
            "Consignee_norm": ["CLIENTE FICTICIO"],
        }))
        self.assertEqual(sheet.max_column, 2)
        self.assertEqual(sheet["B5"].value, datetime(2025, 8, 9))
        self.assertEqual(sheet["B5"].number_format, "DD/MM/YYYY")
        self.assertEqual(len(sheet.tables), 1)

    def test_database_delete_cleans_aliases_and_lists(self):
        db = ClientDatabase(str(self.folder / "clients.db"))
        self.addCleanup(db.close)
        db.add_client("Cliente ficticio", "test@example.invalid")
        client_id = db.get_client_id_by_name("Cliente ficticio")
        db.add_alias(client_id, "Alias ficticio")
        list_id = db.get_all_lists()[0][0]
        db.add_client_to_list(client_id, list_id)
        db.delete_client(client_id)
        self.assertIsNone(db.get_client_by_id(client_id))
        self.assertEqual(db.get_aliases_for_client(client_id), [])
        self.assertEqual(db.get_clients_by_list(list_id), [])

    def test_configuration_round_trip(self):
        with patch.object(ConfigManager, "CONFIG_FILE", str(self.folder / "config.json")):
            config = ConfigManager()
            config.set_email("test@example.invalid")
            config.set_email_body_template("Reporte de prueba: {CLIENT_NAME}")
            reloaded = ConfigManager()
            self.assertEqual(reloaded.get_email(), "test@example.invalid")
            self.assertEqual(reloaded.get_email_body_template(),
                             "Reporte de prueba: {CLIENT_NAME}")

    def test_email_payload_simulated_only(self):
        sender, attachment = self.make_sender()
        with patch("email_service.email_sender.requests.post") as post:
            post.return_value.status_code = 202
            ok, _ = sender.send_email(
                "a@example.invalid, b@example.invalid", "Prueba", "<p>Prueba</p>",
                attachment, cc_emails="c@example.invalid", bcc_emails="d@example.invalid",
            )
        self.assertTrue(ok)
        post.assert_called_once()
        message = post.call_args.kwargs["json"]["message"]
        self.assertEqual(len(message["toRecipients"]), 2)
        self.assertEqual(len(message["ccRecipients"]), 1)
        self.assertEqual(len(message["bccRecipients"]), 1)
        self.assertEqual(base64.b64decode(message["attachments"][0]["ContentBytes"]),
                         b"reporte ficticio")

    def test_email_without_token_does_not_call_transport(self):
        sender, attachment = self.make_sender(token=None)
        with patch("email_service.email_sender.requests.post") as post:
            ok, _ = sender.send_email("test@example.invalid", "Prueba", "Prueba", attachment)
        self.assertFalse(ok)
        post.assert_not_called()

    def test_email_rejection_is_not_success(self):
        sender, attachment = self.make_sender()
        with patch("email_service.email_sender.requests.post") as post:
            post.return_value.status_code = 403
            post.return_value.text = "Rechazo simulado"
            ok, detail = sender.send_email("test@example.invalid", "Prueba", "Prueba", attachment)
        self.assertFalse(ok)
        self.assertIn("Rechazo simulado", detail)

    @unittest.expectedFailure
    def test_known_bug_template_table_is_replaced(self):
        sheet = self.export_report(pd.DataFrame({"Consignee": ["Cliente ficticio"]}),
                                   with_table=True)
        self.assertEqual(len(sheet.tables), 1, "La tabla antigua permanece junto a la nueva")

    @unittest.expectedFailure
    def test_known_bug_date_column_preserves_pending_text(self):
        sheet = self.export_report(pd.DataFrame({"ETA": ["9/8/25", "10/8/25", "PENDIENTE"]}))
        self.assertEqual(sheet["A7"].value, "PENDIENTE")

    @unittest.expectedFailure
    def test_known_bug_macro_error_closes_excel(self):
        processor = FileProcessor.__new__(FileProcessor)
        processor.local_file_path = str(self.folder / "fake.xlsm")
        excel = Mock()
        excel.Run.side_effect = RuntimeError("Fallo simulado de macro")
        with patch("processing.file_processor.win32com.client.DispatchEx", return_value=excel), \
                patch("processing.file_processor.time.sleep"):
            with self.assertRaisesRegex(Exception, "Fallo simulado"):
                processor.run_macro()
        excel.Quit.assert_called_once()

    @unittest.expectedFailure
    def test_known_bug_update_date_without_microseconds(self):
        app = Mock()
        app.get_sharepoint_access_token.return_value = "fake-token"
        processor = Mock()
        processor.get_last_update_date.return_value = "2026-10-04 12:30:00"
        with patch("gui.app.os.path.exists", return_value=False), \
                patch("gui.app.FileProcessor", return_value=processor), \
                patch("gui.app.messagebox.askyesno", return_value=False):
            Application.filter_and_generate(app)
        errors = [call for call in app.update_progress.call_args_list
                  if call.kwargs.get("style") == "Error.TLabel"]
        self.assertEqual(errors, [], "Una fecha valida se interpreta como error")

    @unittest.expectedFailure
    def test_known_bug_generation_failure_stops_send_pipeline(self):
        app = Mock()
        app.get_sharepoint_access_token.return_value = "fake-token"
        app.generate_reports_simplified.side_effect = (
            lambda: Application.generate_reports_simplified(app)
        )
        processor = Mock()
        processor.run_macro.side_effect = RuntimeError("Fallo simulado de macro")
        with patch("gui.app.os.path.exists", return_value=False), \
                patch("gui.app.FileProcessor", return_value=processor):
            Application.filter_generate_and_send(app)
        # send_mails_filtered es un Mock: nunca se envia un correo real.
        app.send_mails_filtered.assert_not_called()


if __name__ == "__main__":
    unittest.main()
